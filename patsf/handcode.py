"""Hand coding: an Excel workbook per patent, filled in by a person and stored like a model's
structure (tag "human"), so that `agree` can compare it with the models.

  patsf handcode US10799093          writes handcode/US10799093B2.xlsx
  patsf import handcode/US10799093B2.xlsx

See handcode folder for examples xlsx files. 
"""

from __future__ import annotations

import re
from pathlib import Path

from . import db, prompts
from .extract import _store, merge_duplicates, resolve_components
from .schema import check_representation

SHEETS = {
    "components": ["id", "name", "parent", "claims"],
    "functionalities": ["id", "component", "text", "target", "claims"],
}
EXAMPLE = {   # one example row per sheet, in grey, to be overwritten
    "components": ["C1", "pump", "C0", "1, 8"],
    "functionalities": ["FN1", "C1", "drive a wheel", "C2", "1, 3"],
}


def template(conn, patent_id: str, folder: Path = Path("handcode")) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    doc, claims = db.document(conn, patent_id), db.claims(conn, patent_id)
    if doc is None or not claims:
        raise LookupError(f"{patent_id}: no stored claims; run fetch first")
    path = folder / f"{patent_id}.xlsx"
    if path.exists():
        raise FileExistsError(f"{path} exists; it is not overwritten")
    folder.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    guide = wb.active
    guide.title = "guide"
    guide.column_dimensions["A"].width = 110
    guide.append([f"Patent {patent_id}: {doc['title'] or ''}"])
    guide.append(["Code the claims as the models do (instructions below), in the sheets components and "
                  "functionalities; replace the grey example rows. Leave out quotes and status."])
    guide.append([""])
    for line in prompts.EXTRACT_SYSTEM.splitlines():
        guide.append([line])
    sheet = wb.create_sheet("claims")
    sheet.append(["number", "depends on", "text"])
    for c in claims:
        text = re.sub(r"^\s*\d+\s*[.)]\s*", "", c["text"]).strip()
        sheet.append([c["number"], ", ".join(map(str, c["parents"])) or "independent", text])
    sheet.column_dimensions["C"].width = 150
    for row in sheet.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
        lines = sum(len(line) // 140 + 1 for line in str(row[2].value or "").splitlines())
        sheet.row_dimensions[row[0].row].height = min(400, 15 * lines)
    sheet.freeze_panes = "C2"
    for name, cols in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(cols)
        ws.append(EXAMPLE[name])
        for cell in ws[2]:
            cell.font = Font(color="999999", italic=True)
        for col, width in zip("ABCDE", [6, 30, 35, 20, 12]):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
    wb.save(path)
    return path


def _claims(value) -> list[int]:
    out: list[int] = []
    for part in re.split(r"[,;\s]+", str(value or "").strip()):
        if m := re.fullmatch(r"(\d+)\s*[-–]\s*(\d+)", part):
            out += list(range(int(m.group(1)), int(m.group(2)) + 1))
        elif part.isdigit():
            out.append(int(part))
    return sorted(set(out))


def load(path: Path) -> tuple[str, dict]:
    """Read a filled-in workbook. Returns (patent id, structure)."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    patent_id = Path(path).stem
    rows = {}
    for name, cols in SHEETS.items():
        out = []
        for values in wb[name].iter_rows(min_row=2, values_only=True):
            r = dict(zip(cols, values))
            if not any(v not in (None, "") for v in values) or list(values[:len(cols)]) == EXAMPLE[name]:
                continue
            out.append(r)
        rows[name] = out
    prov = lambda r: {"claims": _claims(r.get("claims")), "quote": "", "status": "stated"}
    comps = [{"id": str(r["id"] or f"C{i}").strip(), "name": str(r["name"] or "").strip(),
              "parent": (str(r["parent"]).strip() or None) if r.get("parent") not in (None, "") else None,
              "provenance": prov(r)}
             for i, r in enumerate(rows["components"])]
    funcs = [{"id": str(r["id"] or f"FN{i}").strip(), "component": str(r["component"] or "").strip(),
              "text": str(r["text"] or "").strip(),
              "target": str(r["target"]).strip() if r.get("target") not in (None, "") else None,
              "provenance": prov(r)}
             for i, r in enumerate(rows["functionalities"], 1)]
    return patent_id, merge_duplicates(resolve_components({"components": comps, "functionalities": funcs}))


def store(conn, path: Path, tag: str = "human") -> tuple[str, int, list[str]]:
    patent_id, rep = load(path)
    claims = db.claims(conn, patent_id)
    # Only the first claim family is stored (see fetch): other claim numbers are dropped, and so
    # are entries that cite none of the stored claims.
    stored = {c["number"] for c in claims}
    dropped = []
    for kind in ("components", "functionalities"):
        kept = []
        for el in rep[kind]:
            el["provenance"]["claims"] = [k for k in el["provenance"]["claims"] if k in stored]
            (kept if el["provenance"]["claims"] else dropped).append(el)
        rep[kind] = kept
    warnings = check_representation(rep, claims)
    if dropped:
        warnings.append(f"left out {len(dropped)} entries citing only claims outside the first claim family: "
                        + ", ".join(e["id"] for e in dropped))
    if any(w.startswith("schema:") for w in warnings):
        raise ValueError(f"{path}: " + "; ".join(warnings[:5]))
    version = _store(conn, patent_id, tag, "human", "handcode", rep, warnings)
    return patent_id, version, warnings
